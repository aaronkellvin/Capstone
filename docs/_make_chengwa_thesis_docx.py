"""Word copy of docs/chengwaThesis.md (Chapters I–III aligned)."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = Path(__file__).resolve().parent / "chengwaThesis.md"
OUT = Path(__file__).resolve().parent / "chengwaThesis.docx"


def set_run(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "Times New Roman")


def add_runs(paragraph, text, *, size=12, italic_all=False):
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run(run, size=size, bold=True, italic=italic_all)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            set_run(run, size=size, italic=True)
        else:
            run = paragraph.add_run(part)
            set_run(run, size=size, italic=italic_all)


def add_p(doc, text, *, size=12, bold=False, center=False, italic=False, after=8, before=0, justify=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing = 2.0
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold and not re.search(r"\*\*", text):
        run = p.add_run(text)
        set_run(run, size=size, bold=True, italic=italic)
    else:
        add_runs(p, text, size=size, italic_all=italic)
    return p


def shade_header(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(
        qn("w:shd"),
        {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "1F4E79"},
    )
    tcPr.append(shd)


def set_cell(cell, text, *, header=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(text.strip())
    set_run(run, size=9, bold=header)
    if header:
        run.font.color.rgb = RGBColor(255, 255, 255)
        shade_header(cell)


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            set_cell(table.rows[r].cells[c], val, header=(r == 0))
    doc.add_paragraph()


def is_table_sep(line: str) -> bool:
    return bool(re.match(r"^\|[\s:-|]+\|$", line.strip()))


def parse_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if stripped.startswith("|") and i + 1 < len(lines) and is_table_sep(lines[i + 1].strip()):
            rows = [parse_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        if not stripped or stripped == "---":
            i += 1
            continue
        if stripped.startswith("# ") and not stripped.startswith("##"):
            add_p(doc, stripped[2:].strip(), size=16, bold=True, center=True, after=6, before=18)
        elif stripped.startswith("## "):
            add_p(doc, stripped[3:].strip(), size=14, bold=True, after=8, before=16)
        elif stripped.startswith("### "):
            add_p(doc, stripped[4:].strip(), size=13, bold=True, after=6, before=12)
        else:
            add_p(doc, stripped, size=12, justify=True, after=8)
        i += 1

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
