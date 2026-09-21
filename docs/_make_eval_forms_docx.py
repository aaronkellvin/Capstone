"""Printable Word copies of the Bloom teacher and student evaluation forms."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_DIR = Path(__file__).resolve().parent

TEACHER_ITEMS_25010 = [
    ("Functional suitability", "Bloom does the teaching tasks I need (materials, HOTS items, assessments, results)."),
    ("Performance efficiency", "Pages and AI generation finish in a reasonable time."),
    ("Compatibility", "Bloom works in the browser and with my lesson files (for example PDF or Word)."),
    ("Interaction capability", "The screens are easy to understand and use."),
    ("Reliability", "Bloom works consistently without crashing or losing my work."),
    ("Security", "My account and class data feel protected (login, teacher-only tools)."),
    ("Maintainability", "The system is organized; problems are easy to notice or report."),
    ("Flexibility", "I can use Bloom on school computers and for my subject / class size."),
    ("Safety", "I can review AI content before students see it, so wrong or unsafe items are not published."),
]
TEACHER_ITEMS_25019 = [
    ("Beneficialness", "Using Bloom helps me prepare HOTS assessments and support student learning."),
    ("Freedom from risk", "Using Bloom does not put students or the school at unnecessary risk (privacy, wrong items, wasted time)."),
    ("Acceptability", "I trust Bloom enough to use it in class, and it fits school rules."),
]
STUDENT_ITEMS_25010 = [
    ("Functional suitability", "Bloom lets me read summaries, practice questions, take assessments, and see my results."),
    ("Performance efficiency", "Pages and questions load fast enough for class."),
    ("Compatibility", "Bloom works on the computer or device I used."),
    ("Interaction capability", "The screens are easy to understand and use."),
    ("Reliability", "Bloom works without crashing or losing my answers."),
    ("Security", "Only I can use my account, and my scores feel private."),
    ("Maintainability", "If something goes wrong, I can tell what happened or try again."),
    ("Flexibility", "I can use Bloom in school for different lessons or subjects."),
    ("Safety", "Bloom does not show me harmful content, and assessments come from my teacher."),
]
STUDENT_ITEMS_25019 = [
    ("Beneficialness", "Bloom helps me practice higher-order thinking and understand the lesson."),
    ("Freedom from risk", "Using Bloom feels safe (my data, my grades, and the questions)."),
    ("Acceptability", "I am comfortable using Bloom, and I trust what it shows me."),
]


def set_run(run, size=11, bold=False, italic=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")


def add_p(doc, text="", *, size=11, bold=False, italic=False, center=False, after=6, before=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        run = p.add_run(text)
        set_run(run, size=size, bold=bold, italic=italic)
    return p


def shade_header(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(
        qn("w:shd"),
        {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "1F4E79"},
    )
    tcPr.append(shd)


def set_cell(cell, text, *, bold=False, size=10, center=False, header=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run(run, size=size, bold=bold)
    if header:
        run.font.color.rgb = RGBColor(255, 255, 255)
        shade_header(cell)


def add_rating_table(doc, items, start_n=1):
    table = doc.add_table(rows=1 + len(items), cols=8)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["#", "ISO characteristic", "Statement", "1", "2", "3", "4", "5"]
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, bold=True, size=9, center=True, header=True)
    for r, (char, stmt) in enumerate(items, start=1):
        n = start_n + r - 1
        set_cell(table.rows[r].cells[0], str(n), size=9, center=True)
        set_cell(table.rows[r].cells[1], char, size=9, bold=True)
        set_cell(table.rows[r].cells[2], stmt, size=9)
        for c in range(3, 8):
            set_cell(table.rows[r].cells[c], "☐", size=12, center=True)
    widths = [Cm(0.8), Cm(3.4), Cm(8.4), Cm(0.8), Cm(0.8), Cm(0.8), Cm(0.8), Cm(0.8)]
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w
    return len(items)


def build_form(*, title, audience, next_form, respondent_lines, intro_25010, items_25010, items_25019, out_path: Path):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    add_p(doc, "Colegio de San Juan de Letran Calamba — Junior High School", size=10, center=True, after=0)
    add_p(doc, title, size=16, bold=True, center=True, after=2, before=4)
    add_p(doc, "Bloom (AI-Powered Smart Study Assistant)", size=11, center=True, after=2)
    add_p(
        doc,
        f"Rate system quality only. Usability and experience are on the separate {next_form}. Aligned with ISO/IEC 25040:2024, ISO/IEC 25010:2023, and ISO/IEC 25019:2023. All characteristics of both models are included (one item each).",
        size=9,
        italic=True,
        center=True,
        after=8,
    )

    add_p(doc, "A. Respondent  (ISO/IEC 25040)", size=12, bold=True, before=4, after=4)
    for line in respondent_lines:
        add_p(doc, line, size=11, after=4)

    add_p(doc, "B. How to rate", size=12, bold=True, before=8, after=4)
    add_p(doc, "Tick one box for each statement.", size=11, after=4)
    add_p(
        doc,
        "1 = Strongly disagree    2 = Disagree    3 = Neutral    4 = Agree    5 = Strongly agree",
        size=10,
        after=8,
    )

    add_p(doc, "C. Product quality  —  ISO/IEC 25010:2023", size=12, bold=True, before=4, after=2)
    add_p(doc, intro_25010, size=10, italic=True, after=6)
    n = add_rating_table(doc, items_25010, start_n=1)

    add_p(doc, "D. Quality in use  —  ISO/IEC 25019:2023", size=12, bold=True, before=12, after=2)
    add_p(doc, f"Rate using Bloom as a {audience}.", size=10, italic=True, after=6)
    add_rating_table(doc, items_25019, start_n=n + 1)

    add_p(doc, f"Thank you. Please complete the {next_form} next.", size=11, italic=True, before=12)
    try:
        doc.save(out_path)
    except PermissionError:
        fallback = out_path.with_name(out_path.stem + "-updated" + out_path.suffix)
        doc.save(fallback)
        print(f"locked {out_path.name}; wrote {fallback}")
        return
    print(f"wrote {out_path}")


def main():
    build_form(
        title="Teacher Evaluation Form",
        audience="teacher",
        respondent_lines=[
            "Name (optional): ________________________________    Date: ____________________",
            "Subject taught:    ☐ English      ☐ Mathematics      ☐ Science      ☐ Other: __________",
            "How long did you use Bloom?    ☐ Under 30 min      ☐ 30–60 min      ☐ More than 1 hour",
        ],
        next_form="Teacher Questionnaire",
        intro_25010="After using Bloom (upload lessons, generate HOTS items, publish assessments, monitor results), rate the system.",
        items_25010=TEACHER_ITEMS_25010,
        items_25019=TEACHER_ITEMS_25019,
        out_path=OUT_DIR / "Bloom-Teacher-Evaluation-Form.docx",
    )
    build_form(
        title="Student Evaluation Form",
        audience="student",
        respondent_lines=[
            "Name (optional): ________________________________    Date: ____________________",
            "Grade / section: ________________    Subject used:    ☐ English      ☐ Mathematics      ☐ Science",
            "How long did you use Bloom?    ☐ Under 30 min      ☐ 30–60 min      ☐ More than 1 hour",
        ],
        next_form="Student Questionnaire",
        intro_25010="After using Bloom (summaries, practice, assessments, results), rate the system.",
        items_25010=STUDENT_ITEMS_25010,
        items_25019=STUDENT_ITEMS_25019,
        out_path=OUT_DIR / "Bloom-Student-Evaluation-Form.docx",
    )


if __name__ == "__main__":
    main()
