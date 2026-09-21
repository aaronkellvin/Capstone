"""Printable Word copies of the Bloom teacher and student questionnaires."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_DIR = Path(__file__).resolve().parent

TEACHER_SUS = [
    "I would like to use Bloom often.",
    "Bloom was too complicated.",
    "Bloom was easy to use.",
    "I needed help from a technical person to use Bloom.",
    "The parts of Bloom worked well together.",
    "Things in Bloom felt inconsistent.",
    "Most teachers could learn Bloom quickly.",
    "Bloom was awkward or tiring to use.",
    "I felt confident using Bloom.",
    "I had to learn too many things before I could use Bloom.",
]
STUDENT_SUS = [
    "I would like to use Bloom often.",
    "Bloom was too complicated.",
    "Bloom was easy to use.",
    "I needed help from someone to use Bloom.",
    "The parts of Bloom worked well together.",
    "Things in Bloom felt inconsistent.",
    "Most students could learn Bloom quickly.",
    "Bloom was awkward or tiring to use.",
    "I felt confident using Bloom.",
    "I had to learn too many things before I could use Bloom.",
]


def set_run(run, size=11, bold=False, italic=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")


def add_p(doc, text="", *, size=11, bold=False, italic=False, center=False, after=6, before=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        run = p.add_run(text)
        set_run(run, size=size, bold=bold, italic=italic)
    return p


def shade_header(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(
        qn("w:shd"),
        {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "1F4E79"},
    )
    tcPr.append(shd)


def set_cell(cell, text, *, bold=False, size=10, center=False, header=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run(run, size=size, bold=bold)
    if header:
        run.font.color.rgb = RGBColor(255, 255, 255)
        shade_header(cell)


def add_sus_table(doc, items, start_n=1):
    table = doc.add_table(rows=1 + len(items), cols=7)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["#", "Statement", "1", "2", "3", "4", "5"]
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, bold=True, size=9, center=True, header=True)
    for r, stmt in enumerate(items, start=1):
        n = start_n + r - 1
        set_cell(table.rows[r].cells[0], str(n), size=9, center=True)
        set_cell(table.rows[r].cells[1], stmt, size=9)
        for c in range(2, 7):
            set_cell(table.rows[r].cells[c], "☐", size=12, center=True)
    widths = [Cm(0.8), Cm(11.8), Cm(0.8), Cm(0.8), Cm(0.8), Cm(0.8), Cm(0.8)]
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w


def add_open(doc, number, prompt):
    add_p(doc, f"{number}. {prompt}", size=11, bold=True, before=8, after=4)
    for _ in range(2):
        add_p(doc, "_" * 92, size=11, after=8)


def build(*, title, eval_form_name, profile_lines, sus_items, open_items, out_path: Path):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    add_p(doc, "Colegio de San Juan de Letran Calamba — Junior High School", size=10, center=True, after=0)
    add_p(doc, title, size=16, bold=True, center=True, after=2, before=4)
    add_p(doc, "Bloom (AI-Powered Smart Study Assistant)", size=11, center=True, after=2)
    add_p(
        doc,
        "Researchers: Esguerra, Patrick Carlos P.; Lamadrid, Isaiah Matthew C.; Quilitis, Marvic Mat M.",
        size=9,
        center=True,
        after=6,
    )
    add_p(
        doc,
        f"This questionnaire is for usability and experience after you try Bloom. System quality (ISO/IEC 25010:2023 and ISO/IEC 25019:2023) is on the separate {eval_form_name}. Your answers are for this capstone study only. You may skip your name. Completing this form means you agree to take part.",
        size=9,
        italic=True,
        after=10,
    )

    add_p(doc, "Part I. Respondent profile", size=13, bold=True, after=6)
    for line in profile_lines:
        add_p(doc, line, size=11, after=5)
    add_p(doc, "Date: ____________________", size=11, after=10)

    add_p(doc, "How to answer Part II", size=13, bold=True, before=4, after=4)
    add_p(doc, "Tick one box for each statement.", size=11, after=2)
    add_p(
        doc,
        "1 = Strongly disagree    2 = Disagree    3 = Neutral    4 = Agree    5 = Strongly agree",
        size=10,
        after=10,
    )

    add_p(doc, "Part II. Usability  —  adapted SUS", size=13, bold=True, after=6)
    add_sus_table(doc, sus_items, start_n=1)

    add_p(doc, "Part III. Open questions", size=13, bold=True, before=12, after=4)
    for number, prompt in open_items:
        add_open(doc, number, prompt)

    add_p(doc, "Thank you.", size=11, italic=True, before=6)
    doc.save(out_path)
    print(f"wrote {out_path}")


def main():
    build(
        title="Teacher Questionnaire",
        eval_form_name="Teacher Evaluation Form",
        profile_lines=[
            "1. Name (optional): ________________________________",
            "2. Sex:    ☐ Female      ☐ Male      ☐ Prefer not to say",
            "3. Years of teaching:    ☐ 1–5      ☐ 6–10      ☐ 11–15      ☐ 16 or more",
            "4. Subject:    ☐ English      ☐ Mathematics      ☐ Science      ☐ Other: __________",
            "5. Device used:    ☐ School computer      ☐ Laptop      ☐ Tablet      ☐ Phone",
            "6. How long did you use Bloom today?    ☐ Under 30 min      ☐ 30–60 min      ☐ More than 1 hour",
        ],
        sus_items=TEACHER_SUS,
        open_items=[
            (11, "What is hard about making HOTS assessments without Bloom?"),
            (12, "How did Bloom help (or not help) with curriculum-aligned HOTS items?"),
            (13, "What should we improve?"),
        ],
        out_path=OUT_DIR / "Bloom-Teacher-Questionnaire.docx",
    )
    build(
        title="Student Questionnaire",
        eval_form_name="Student Evaluation Form",
        profile_lines=[
            "1. Name (optional): ________________________________",
            "2. Sex:    ☐ Female      ☐ Male      ☐ Prefer not to say",
            "3. Grade / section: ________________________________",
            "4. Subject used today:    ☐ English      ☐ Mathematics      ☐ Science",
            "5. Device used:    ☐ School computer      ☐ Laptop      ☐ Tablet      ☐ Phone",
            "6. How long did you use Bloom today?    ☐ Under 30 min      ☐ 30–60 min      ☐ More than 1 hour",
        ],
        sus_items=STUDENT_SUS,
        open_items=[
            (11, "What was hard about practicing higher-order thinking before Bloom?"),
            (12, "How did Bloom help you understand or practice the lesson?"),
            (13, "What was hard or confusing? What should we improve?"),
        ],
        out_path=OUT_DIR / "Bloom-Student-Questionnaire.docx",
    )


if __name__ == "__main__":
    main()
